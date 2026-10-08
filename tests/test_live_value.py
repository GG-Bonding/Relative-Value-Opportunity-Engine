"""The live fair value uses the rate spread. It does not fill a missing day with zero."""

from __future__ import annotations

import math
from datetime import date, datetime, timedelta
from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook

from domain.config import EngineConfig
from domain.timeutil import UTC
from forward.live_value import rate_fair_value
from forward.official_rates import parse_boe_nominal_archive

AS_OF = datetime(2026, 10, 7, tzinfo=UTC)


def test_archive_reads_the_nominal_spot_sheet_and_skips_old_files() -> None:
    payload = _archive()
    points = parse_boe_nominal_archive(payload)
    assert points[date(2016, 1, 4)] == 4.2
    assert points[date(2025, 1, 2)] == 4.4
    assert date(1979, 1, 2) not in points


def test_rate_fair_value_ignores_the_session_it_is_pricing() -> None:
    closes, uk, de = _history(280)
    last = max(closes)
    closes[last] = 5.0
    fitted = rate_fair_value(closes, uk, de, AS_OF, EngineConfig())
    assert fitted is not None
    assert fitted.day == last
    assert fitted.sessions >= 252
    assert fitted.price < 1.0


def test_a_short_history_stays_empty() -> None:
    closes, uk, de = _history(100)
    assert rate_fair_value(closes, uk, de, AS_OF, EngineConfig()) is None


def test_a_one_sided_curve_is_not_filled() -> None:
    closes, uk, de = _history(280)
    missing = max(closes)
    del uk[missing]
    fitted = rate_fair_value(closes, uk, de, AS_OF, EngineConfig())
    assert fitted is not None
    assert fitted.day != missing


def _history(sessions: int) -> tuple[dict[date, float], dict[date, float], dict[date, float]]:
    closes: dict[date, float] = {}
    uk: dict[date, float] = {}
    de: dict[date, float] = {}
    day = date(2024, 1, 2)
    index = 0
    while len(closes) < sessions:
        if day.weekday() < 5:
            level = 1.2 + index * 0.001
            uk[day] = 4.0 + level
            de[day] = 4.0
            closes[day] = math.exp(-0.01 * level)
            index += 1
        day += timedelta(days=1)
    return closes, uk, de


def _archive() -> bytes:
    packed = BytesIO()
    with ZipFile(packed, "w", compression=ZIP_DEFLATED) as archive:
        archive.writestr("GLC Nominal daily data_1979 to 1984.xlsx", _book("1979-01-02", 9.99))
        archive.writestr("GLC Nominal daily data_2016 to 2024.xlsx", _book("2016-01-04", 4.2))
        archive.writestr("GLC Nominal daily data_2025 to present.xlsx", _book("2025-01-02", 4.4))
    return packed.getvalue()


def _book(day: str, spot: float) -> bytes:
    book = Workbook()
    forward = book.active
    assert forward is not None
    forward.title = "2. nominal fwd curve"
    _tenors(forward)
    forward["A6"] = datetime.fromisoformat(day)
    forward["E6"] = 9.99
    curve = book.create_sheet("4. nominal spot curve")
    _tenors(curve)
    curve["A6"] = datetime.fromisoformat(day)
    curve["E6"] = spot
    sheet = BytesIO()
    book.save(sheet)
    book.close()
    return sheet.getvalue()


def _tenors(sheet: object) -> None:
    from openpyxl.worksheet.worksheet import Worksheet

    assert isinstance(sheet, Worksheet)
    sheet["A4"] = "years"
    for column, tenor in enumerate((0.5, 1, 1.5, 2, 2.5), start=2):
        sheet.cell(4, column, tenor)
