"""One automatic EURGBP round. The three layers stay separate, and the log only grows."""

from __future__ import annotations

import json
import sys
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import Workbook

from api.cli import main
from data.store import PitStore
from domain.config import EngineConfig
from domain.enums import ExecutionMode
from domain.timeutil import UTC
from forward.mt5 import Mt5Tick
from forward.official_rates import (
    _as_day,
    align_same_day,
    parse_boe_nominal_spot,
    parse_bundesbank_de2y,
)
from forward.prints import InformationPrint, normalize_print, read_information, usable_prints
from forward.rates import RateSnapshot
from forward.round import RoundInputs, run_round
from forward.terminal import read_eurgbp_tick

ENTRY = datetime(2026, 10, 1, 10, tzinfo=UTC)
DAY_2 = datetime(2026, 10, 2, 10, tzinfo=UTC)
FAIR = 0.8515
CONFIG = Path(__file__).resolve().parents[1] / "configs" / "eurgbp.toml"


def test_spot_curve_uses_the_two_year_tenor_and_skips_blanks() -> None:
    points = parse_boe_nominal_spot(_nominal_zip())
    assert points[date(2026, 10, 1)] == pytest.approx(4.2)
    assert date(2026, 10, 2) not in points
    assert 9.99 not in points.values()
    assert 0 not in points.values()
    assert points[date(2024, 1, 1)] == pytest.approx(4.4)


def test_bundesbank_dot_is_missing_not_zero() -> None:
    points = parse_bundesbank_de2y(
        "\n".join(
            [
                "unit,Prozent,",
                "1997-08-01,.,No value available",
                "2026-10-01,2.10,",
                "2026-10-02,.,No value available",
                "2026-10-03,2.40,",
            ]
        )
    )
    assert points[date(2026, 10, 1)] == pytest.approx(2.10)
    assert date(2026, 10, 2) not in points
    assert 0 not in points.values()


def test_spread_uses_only_the_shared_trading_day() -> None:
    uk = {date(2026, 10, 1): 4.2, date(2026, 10, 2): 4.5}
    de = {date(2026, 10, 1): 2.1}
    aligned = align_same_day(uk, de, datetime(2026, 10, 3, 12, tzinfo=UTC))
    assert aligned.uk2y == pytest.approx(4.2)
    assert aligned.de2y == pytest.approx(2.1)
    assert aligned.rate_diff == pytest.approx(2.1)
    assert aligned.uk2y_time is not None and aligned.de2y_time is not None
    assert aligned.uk2y_time.date() == aligned.de2y_time.date()
    early = align_same_day(uk, de, datetime(2026, 10, 2, 12, tzinfo=UTC))
    assert early.uk2y == pytest.approx(4.2)
    mismatched = align_same_day(
        {date(2026, 10, 2): 4.5},
        {date(2026, 10, 1): 2.1},
        datetime(2026, 10, 3, 12, tzinfo=UTC),
    )
    assert mismatched.rate_diff is None
    assert mismatched.uk2y == pytest.approx(4.5)
    assert mismatched.de2y == pytest.approx(2.1)
    assert 0 not in {mismatched.uk2y, mismatched.de2y}


def test_calendar_surprise_is_a_factor_and_a_flash_is_not() -> None:
    feed = _Feed()
    prints, error = read_information(feed)
    assert error is None
    assert feed.quoted is False
    usable = usable_prints(prints, ENTRY)
    calendar = next(item for item in usable if item.kind == "CALENDAR" and item.indicator == "CPI")
    assert calendar.country == "UK"
    assert calendar.actual == pytest.approx(3.2)
    assert calendar.forecast == pytest.approx(3.0)
    assert calendar.previous == pytest.approx(2.8)
    assert calendar.importance == "3"
    assert calendar.surprise == pytest.approx(0.2)
    assert all(item.event_time is None or item.event_time <= ENTRY for item in usable)
    assert not any(item.headline.startswith("欧元区10月") for item in usable)
    flash = next(item for item in usable if item.kind == "FLASH")
    assert flash.surprise is None
    assert flash.indicator is None
    assert flash.country is None
    german = next(item for item in usable if item.headline.startswith("德国"))
    assert german.surprise is None
    assert german.country is None


def test_uk_cpi_surprise_does_not_become_a_short_without_fair_value(tmp_path: Path) -> None:
    result = _round(tmp_path, _inputs(prints=[_cpi()]), None)
    assert result["decision"] == "WATCH"
    assert result["model_readiness"] == "MODEL_NOT_READY"
    assert result["direction"] is None
    assert result["event_pressure"] == pytest.approx(-0.2 / 3)
    assert result["narratives"] == ["UK CPI above expectation → GBP relative support → EURGBP bearish factor"]
    assert "SHORT" not in result["reasons"]


def test_missing_quote_is_data_degraded_and_does_not_use_a_cross(tmp_path: Path) -> None:
    result = _round(tmp_path, _inputs(quote=False, quote_error="MT5 returned no EURGBP tick"), FAIR)
    assert result["decision"] == "DATA_DEGRADED"
    assert result["decision"] not in {"LONG", "SHORT", "NO_TRADE"}
    assert result["inputs"]["quote"] is None
    assert result["inputs"]["rates"]["rate_diff"] == pytest.approx(1.65)
    assert result["executed"] is False
    store = PitStore(tmp_path / "round.duckdb")
    assert store.rows("market_observations") == []
    assert {row["curve_id"] for row in store.rows("rate_observations")} == {"UK2Y", "DE2Y"}
    assert all(row["rate"] != 0 for row in store.rows("rate_observations"))
    store.close()


def test_missing_curve_is_data_degraded_and_not_zero(tmp_path: Path) -> None:
    rates = RateSnapshot(uk2y=4.2, de2y=None, uk2y_time=ENTRY, de2y_time=None)
    result = _round(tmp_path, _inputs(rates=rates), FAIR)
    assert result["decision"] == "DATA_DEGRADED"
    assert result["inputs"]["rates"]["rate_diff"] is None
    assert result["inputs"]["rates"]["de2y"] is None
    store = PitStore(tmp_path / "round.duckdb")
    stored = store.rows("rate_observations")
    assert [row["curve_id"] for row in stored] == ["UK2Y"]
    assert stored[0]["rate"] == pytest.approx(4.2)
    store.close()


def test_shadow_paper_and_live_share_the_signal(tmp_path: Path) -> None:
    book = _inputs(prints=[_cpi()])
    shadow = _round(tmp_path, book, FAIR, mode=ExecutionMode.SHADOW, name="shadow")
    paper = _round(tmp_path, book, FAIR, mode=ExecutionMode.PAPER, name="paper")
    live = _round(tmp_path, book, FAIR, mode=ExecutionMode.LIVE, name="live")
    assert shadow["decision"] == paper["decision"] == live["decision"] == "SHORT"
    assert shadow["reasons"] == paper["reasons"] == live["reasons"]
    assert shadow["fill_price"] is None
    assert shadow["executed"] is False
    assert paper["executed"] is True
    assert paper["fill_price"] == pytest.approx(0.85567)
    assert live["executed"] is False
    assert live["fill_price"] is None
    assert live["live_order"] == "not armed"


def test_later_round_appends_and_leaves_the_opening_row(tmp_path: Path) -> None:
    log = tmp_path / "rounds.jsonl"
    opened = _round(tmp_path, _inputs(prints=[_cpi()]), FAIR, mode=ExecutionMode.PAPER, log=log)
    first = log.read_text(encoding="utf-8")
    store = PitStore(tmp_path / "round.duckdb")
    frozen = _opening(store, opened["opportunity_id"])
    store.close()
    again = _round(tmp_path, _inputs(prints=[_cpi()]), FAIR, mode=ExecutionMode.PAPER, log=log)
    assert again["duplicate"] is True
    moved = _inputs(tick=Mt5Tick("EURGBP", DAY_2, 0.8529, 0.8535))
    followed = _round(tmp_path, moved, 0.8514, mode=ExecutionMode.PAPER, log=log)
    assert followed["status"] == "CONVERGING"
    lines = log.read_text(encoding="utf-8").splitlines()
    assert len(lines) == 3
    assert lines[0] + "\n" == first
    store = PitStore(tmp_path / "round.duckdb")
    still = _opening(store, opened["opportunity_id"])
    assert still["sequence"] == 0
    assert still["mid"] == pytest.approx(float(frozen["mid"]))
    assert still["fill_price"] == pytest.approx(float(frozen["fill_price"]))
    assert still["fair_value"] == pytest.approx(FAIR)
    store.close()


def test_terminal_reader_asks_for_the_eurgbp_tick_and_sends_nothing(monkeypatch: pytest.MonkeyPatch) -> None:
    calls: list[str] = []

    class _Tick:
        bid = 0.8557
        ask = 0.8563
        time = 1_759_327_200
        time_msc = 1_759_327_200_000

    class _Terminal:
        def initialize(self, path: str | None = None) -> bool:
            calls.append(f"initialize:{path}")
            return True

        def symbol_select(self, symbol: str, enable: bool) -> bool:
            calls.append(f"select:{symbol}:{enable}")
            return True

        def symbol_info_tick(self, symbol: str) -> _Tick:
            calls.append(f"tick:{symbol}")
            return _Tick()

        def shutdown(self) -> None:
            calls.append("shutdown")

        def order_send(self, request: object) -> None:
            raise AssertionError(request)

    monkeypatch.setitem(sys.modules, "MetaTrader5", _Terminal())
    monkeypatch.delenv("MT5_TERMINAL", raising=False)
    tick = read_eurgbp_tick(terminal="C:/terminal64.exe")
    assert tick.symbol == "EURGBP"
    assert tick.bid == pytest.approx(0.8557)
    assert tick.ask == pytest.approx(0.8563)
    assert calls == ["initialize:C:/terminal64.exe", "select:EURGBP:True", "tick:EURGBP", "shutdown"]
    source = Path(__file__).resolve().parents[1].joinpath("src/forward/terminal.py").read_text(encoding="utf-8")
    assert "order_send" not in source


def test_excel_serial_is_a_calendar_day() -> None:
    assert _as_day(45292) == date(2024, 1, 1)


def test_cli_round_defaults_to_shadow_and_live_stays_unarmed(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    def fake(terminal: str | None, as_of: datetime) -> RoundInputs:
        assert terminal is None
        assert as_of.tzinfo is not None
        return _inputs()

    monkeypatch.setattr("forward.collect.load_round_inputs", fake)
    log = tmp_path / "rounds.jsonl"
    code = main(
        ["--db", str(tmp_path / "shadow.duckdb"), "--config", str(CONFIG), "round", "--log", str(log)]
    )
    assert code == 0
    recorded = json.loads(log.read_text(encoding="utf-8"))
    assert recorded["mode"] == "SHADOW"
    assert recorded["decision"] == "WATCH"
    assert recorded["model_readiness"] == "MODEL_NOT_READY"
    live_log = tmp_path / "live.jsonl"
    live = main(
        [
            "--db",
            str(tmp_path / "live.duckdb"),
            "--config",
            str(CONFIG),
            "round",
            "--mode",
            "live",
            "--log",
            str(live_log),
        ]
    )
    assert live == 2
    assert json.loads(live_log.read_text(encoding="utf-8"))["live_order"] == "not armed"


def test_normalize_print_reads_the_published_fields() -> None:
    item = normalize_print(
        {
            "title": "英国9月CPI年率",
            "actual": "3.2",
            "consensus": "3.0",
            "previous": "2.9",
            "revised": "2.8",
            "pub_time": "2026-10-01 18:00",
            "star": 3,
        },
        "CALENDAR",
    )
    assert item is not None
    assert item.country == "UK"
    assert item.indicator == "CPI"
    assert item.previous == pytest.approx(2.8)
    assert item.event_time == datetime(2026, 10, 1, 10, tzinfo=UTC)


class _Feed:
    def __init__(self) -> None:
        self.quoted = False

    def quote(self, code: str) -> dict[str, str]:
        self.quoted = True
        raise AssertionError(code)

    def calendar(self) -> list[dict[str, object]]:
        return [
            {
                "title": "英国9月CPI年率",
                "actual": "3.2",
                "consensus": "3.0",
                "previous": "2.9",
                "revised": "2.8",
                "pub_time": "2026-10-01 18:00",
                "star": 3,
            },
            {
                "title": "欧元区10月核心CPI年率初值",
                "actual": None,
                "consensus": "2.2",
                "pub_time": "2026-10-02 17:00",
            },
            {
                "title": "德国10月CPI年率",
                "actual": "2.4",
                "consensus": "2.1",
                "pub_time": "2026-10-01 14:00",
            },
        ]

    def flash(self) -> list[dict[str, object]]:
        return [{"title": "英国央行意外 BUY EURGBP", "importance": "高"}]

    def news(self) -> list[dict[str, object]]:
        return [{"title": "欧元区官员讲话"}]


def _round(
    tmp_path: Path,
    inputs: RoundInputs,
    fair_value: float | None,
    *,
    mode: ExecutionMode = ExecutionMode.SHADOW,
    name: str = "round",
    log: Path | None = None,
) -> dict[str, object]:
    store = PitStore(tmp_path / f"{name}.duckdb")
    try:
        return run_round(
            store,
            EngineConfig(),
            mode,
            inputs,
            ingested_at=datetime(2026, 10, 5, 12, tzinfo=UTC),
            fair_value=fair_value,
            log_path=tmp_path / "rounds.jsonl" if log is None else log,
        )
    finally:
        store.close()


def _inputs(
    *,
    prints: list[InformationPrint] | None = None,
    rates: RateSnapshot | None = None,
    tick: Mt5Tick | None = None,
    quote: bool = True,
    quote_error: str | None = None,
) -> RoundInputs:
    book = tick
    if quote and book is None:
        book = Mt5Tick("EURGBP", ENTRY, 0.8557, 0.8563)
    return RoundInputs(
        prints=[] if prints is None else prints,
        information_error=None,
        rates=RateSnapshot(4.2, 2.55, ENTRY, ENTRY) if rates is None else rates,
        rate_error=None,
        tick=book,
        quote_error=quote_error,
    )


def _cpi() -> InformationPrint:
    return InformationPrint("CALENDAR", ENTRY, "UK", "CPI", 3.2, 3.0, 2.9, "3", "英国9月CPI年率")


def _opening(store: PitStore, opportunity_id: object) -> dict[str, object]:
    rows = [
        row
        for row in store.rows("opportunity_journal")
        if row["opportunity_id"] == opportunity_id and row["sequence"] == 0
    ]
    assert len(rows) == 1
    return rows[0]


def _nominal_zip() -> bytes:
    book = Workbook()
    forward = book.active
    assert forward is not None
    forward.title = "2. fwd curve"
    _tenors(forward)
    forward["A6"] = datetime(2026, 10, 1)
    forward["E6"] = 9.99
    spot = book.create_sheet("4. spot curve")
    _tenors(spot)
    spot["A6"] = datetime(2026, 10, 1)
    spot["E6"] = 4.2
    spot["A7"] = datetime(2026, 10, 2)
    spot["A8"] = 45292
    spot["E8"] = 4.4
    sheet = BytesIO()
    book.save(sheet)
    book.close()
    packed = BytesIO()
    with ZipFile(packed, "w", compression=ZIP_DEFLATED) as archive:
        archive.writestr("GLC Nominal daily data current month.xlsx", sheet.getvalue())
    return packed.getvalue()


def _tenors(sheet: object) -> None:
    from openpyxl.worksheet.worksheet import Worksheet

    assert isinstance(sheet, Worksheet)
    sheet["A4"] = "years"
    for column, tenor in enumerate((0.5, 1, 1.5, 2, 2.5, 20), start=2):
        sheet.cell(4, column, tenor)
